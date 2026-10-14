import unittest
import sys
import os
import importlib.util
import tempfile

package_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if package_root not in sys.path:
    sys.path.insert(0, package_root)

from prettyppt.templates.registry import list_scenarios, list_styles, get_template, get_all_templates
from prettyppt.core.generator import generate_html_deck, generate_pptx_deck
from prettyppt.core.models import Presentation, Slide, Theme

class TestPrettyPPT(unittest.TestCase):
    def test_scenarios_count(self):
        scenarios = list_scenarios()
        self.assertEqual(len(scenarios), 10)

    def test_styles_count(self):
        core_styles = list_styles(only_core=True)
        self.assertEqual(len(core_styles), 8)
        all_styles = list_styles(only_core=False)
        self.assertEqual(len(all_styles), 16)

    def test_all_80_templates_instantiation(self):
        decks = get_all_templates(only_core=True)
        self.assertEqual(len(decks), 80)
        for deck in decks:
            self.assertEqual(len(deck.slides), 5)
            self.assertTrue(len(deck.title) > 0)
            self.assertIsNotNone(deck.theme)

    def test_html_generation(self):
        deck = get_template("pitch_deck", "apple_titanium")
        html = generate_html_deck(deck)
        self.assertIn("<!DOCTYPE html>", html)
        self.assertIn(deck.title, html)
        self.assertIn(deck.theme.name_zh, html)

    def test_serialization(self):
        deck = get_template("product_launch", "silicon_valley")
        json_str = deck.to_json()
        self.assertIn("Project Nebula", json_str)
        restored = Presentation.from_dict(deck.to_dict())
        self.assertEqual(restored.title, deck.title)
        self.assertEqual(len(restored.slides), 5)

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx is not installed")
    def test_pptx_preserves_card_content(self):
        from pptx import Presentation as PptxPresentation

        deck = get_template("pitch_deck", "apple_titanium")
        with tempfile.TemporaryDirectory() as output_dir:
            output_path = os.path.join(output_dir, "cards.pptx")
            self.assertTrue(generate_pptx_deck(deck, output_path))
            exported = PptxPresentation(output_path)
            self.assertEqual(len(exported.slides), len(deck.slides))
            card_text = "\n".join(
                shape.text for shape in exported.slides[2].shapes if shape.has_text_frame
            )
            for card in deck.slides[2].cards:
                for field in (card.title, card.desc, card.tag, card.metric):
                    if field:
                        self.assertIn(field, card_text)

if __name__ == "__main__":
    unittest.main()
